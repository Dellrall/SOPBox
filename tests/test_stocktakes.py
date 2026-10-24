from conftest import sign_up
from io import BytesIO

from openpyxl import load_workbook

def _count_body(product_id, expected, counted):
    return {
        "count_date": "2026-10-24",
        "notes": "Monthly count",
        "lines": [
            {
                "product_id": product_id,
                "expected_stock": expected,
                "counted_stock": counted,
            }
        ],
    }


def test_stocktake_adjusts_stock_and_keeps_auditable_history(client, catalogue_ids):
    product_id = catalogue_ids["fries"]

    saved = client.post("/api/stocktakes", json=_count_body(product_id, 0, 12))
    assert saved.status_code == 201, saved.text
    detail = saved.json()
    assert detail["reference"] == "ST-000001"
    assert detail["counted_by"] == "Olive Owner"
    assert detail["line_count"] == detail["changed_count"] == 1
    assert detail["lines"][0] == {
        "id": detail["lines"][0]["id"],
        "product_id": product_id,
        "product_name": "Fries tray",
        "base_unit": "pcs",
        "expected_stock": 0,
        "counted_stock": 12,
        "difference": 12,
    }
    product = next(row for row in client.get("/api/products").json() if row["id"] == product_id)
    assert product["stock"] == 12

    history = client.get("/api/stocktakes")
    assert history.status_code == 200
    assert history.json()[0]["reference"] == detail["reference"]
    assert client.get(f"/api/stocktakes/{detail['id']}").json()["lines"][0]["difference"] == 12
    assert any(row["action"] == "stocktake.created" for row in client.get("/api/audit").json())


def test_explicit_zero_is_saved_and_stale_counts_are_rejected(client, catalogue_ids):
    product_id = catalogue_ids["fries"]
    assert client.post("/api/stocktakes", json=_count_body(product_id, 0, 8)).status_code == 201

    zero = client.post("/api/stocktakes", json=_count_body(product_id, 8, 0))
    assert zero.status_code == 201
    assert zero.json()["lines"][0]["counted_stock"] == 0
    assert zero.json()["lines"][0]["difference"] == -8
    assert next(row for row in client.get("/api/products").json() if row["id"] == product_id)["stock"] == 0

    stale = client.post("/api/stocktakes", json=_count_body(product_id, 8, 4))
    assert stale.status_code == 409
    assert "stock changed" in stale.json()["detail"]
    assert len(client.get("/api/stocktakes").json()) == 2


def test_excel_template_is_blind_and_upload_preview_ignores_blank_counts(client, catalogue_ids):
    product_id = catalogue_ids["fries"]
    response = client.get("/api/stocktakes/template.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    assert {"Count sheet", "Instructions"}.issubset(workbook.sheetnames)
    sheet = workbook["Count sheet"]
    assert [cell.value for cell in sheet[1]] == [
        "Product ID",
        "Product Name",
        "Category",
        "Counted Stock",
        "Base Unit",
    ]
    row_number = next(row[0].row for row in sheet.iter_rows(min_row=2) if row[0].value == product_id)
    assert sheet.cell(row_number, 4).value is None
    sheet.cell(row_number, 4, 0)
    uploaded = BytesIO()
    workbook.save(uploaded)

    preview = client.post(
        "/api/stocktakes/preview",
        files={
            "file": (
                "count.xlsx",
                uploaded.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["lines"] == [
        {
            "product_id": product_id,
            "product_name": "Fries tray",
            "base_unit": "pcs",
            "expected_stock": 0,
            "counted_stock": 0,
            "difference": 0,
        }
    ]


def test_staff_section_access_is_enforced_on_stocktake_and_catalogue_apis(
    client, make_user, catalogue_ids
):
    staff = make_user("staff", "counter")
    staff_info = staff.get("/api/auth/me").json()
    assert staff_info["section_access"]["stocktake"] == "view"
    user_id = staff_info["id"]
    access = {
        key: "none"
        for key in (
            "dashboard",
            "receive",
            "records",
            "products",
            "purchasing",
            "approvals",
            "activity",
        )
    }
    access["stocktake"] = "view"
    updated = client.patch(f"/api/users/{user_id}", json={"section_access": access})
    assert updated.status_code == 200, updated.text
    assert staff.get("/api/stocktakes").status_code == 200
    assert staff.get("/api/stocktakes/products").status_code == 200
    assert staff.get("/api/products").status_code == 403

    access["stocktake"] = "edit"
    updated = client.patch(f"/api/users/{user_id}", json={"section_access": access})
    assert updated.status_code == 200, updated.text

    saved = staff.post(
        "/api/stocktakes",
        json=_count_body(catalogue_ids["fries"], 0, 3),
    )
    assert saved.status_code == 201, saved.text


def test_staff_cannot_edit_products_with_view_only_access(client, make_user, catalogue_ids):
    staff = make_user("staff", "viewer")
    assert staff.get("/api/products").status_code == 200
    category_id = catalogue_ids["packaging_category"]
    response = staff.post(
        "/api/products",
        json={
            "name": "Unauthorized product",
            "category_id": category_id,
            "base_unit": "pcs",
            "temp_min_c": None,
            "temp_max_c": None,
            "min_stock": 0,
            "sop": "",
            "pack_units": [],
        },
    )
    assert response.status_code == 403


def test_dashboard_access_does_not_grant_product_or_delivery_records(client, make_user):
    staff = make_user("staff", "dashboardviewer")
    staff_info = staff.get("/api/auth/me").json()
    access = {key: "none" for key in staff_info["section_access"]}
    access["dashboard"] = "view"
    updated = client.patch(
        f"/api/users/{staff_info['id']}", json={"section_access": access}
    )
    assert updated.status_code == 200, updated.text

    assert staff.get("/api/dashboard").status_code == 200
    assert staff.get("/api/products").status_code == 403
    assert staff.get("/api/deliveries").status_code == 403


def test_stocktake_records_are_isolated_between_workspaces(client, catalogue_ids):
    saved = client.post(
        "/api/stocktakes",
        json=_count_body(catalogue_ids["fries"], 0, 7),
    )
    assert saved.status_code == 201

    other_workspace = sign_up(workspace="Other Kitchen", username="other")
    assert other_workspace.get("/api/stocktakes").json() == []
    assert other_workspace.get(f"/api/stocktakes/{saved.json()['id']}").status_code == 404
