"""The sales sheet: laid out like the count sheet, bundles and loose sold, and the lot each sale came from."""

import io

from openpyxl import load_workbook

from test_lots import _day, _lots, _receive


def _upload(client, store, rows):
    workbook = load_workbook(io.BytesIO(client.get(f"/api/sales/template.xlsx?storage_id={store}").content))
    sheet = workbook["Sales sheet"]
    header = next(r for r in range(1, sheet.max_row + 1) if sheet.cell(r, 1).value == "Product ID")
    for product_id, bundles, loose, lot in rows:
        row = next(r for r in range(header, sheet.max_row + 1) if sheet.cell(r, 1).value == product_id)
        if sheet.cell(row, 4).value is not None or sheet.cell(row, 6).value is not None:
            sheet.insert_rows(row + 1)  # a second lot of the same product: a copied row
            row += 1
            sheet.cell(row, 1, product_id)
        sheet.cell(row, 4, bundles)
        sheet.cell(row, 6, loose)
        sheet.cell(row, 8, lot)
    out = io.BytesIO()
    workbook.save(out)
    return client.post(f"/api/sales/preview?storage_id={store}", files={"file": ("sales.xlsx", out.getvalue())})


def test_sales_sheet_takes_sales_off_the_named_lot(client, catalogue_ids):
    milk = client.post(
        "/api/products",
        json={"name": "Milk", "category_id": catalogue_ids["packaging_category"], "base_unit": "btl",
              "track_expiry": True, "pack_units": [{"unit": "crate", "base_qty": 6}]},
    ).json()["id"]
    _receive(client, milk, 10, expiry_on=_day(10))
    _receive(client, milk, 8, expiry_on=_day(20))
    early, late = _lots(client, milk)
    store = early["storage_id"]

    workbook = load_workbook(io.BytesIO(client.get(f"/api/sales/template.xlsx?storage_id={store}").content))
    assert [c.value for c in workbook["Lots"]["A"]][1:] == [early["code"], late["code"]]
    assert "Lot Used" in [c.value for c in next(r for r in workbook["Sales sheet"].iter_rows() if r[0].value == "Product ID")]

    preview = _upload(client, store, [(milk, 1, 1, late["code"]), (milk, None, 2, None)])
    assert preview.status_code == 200, preview.text
    lines = preview.json()["lines"]
    assert [(line["quantity"], line["lot_code"]) for line in lines] == [(2, None), (7, late["code"])]

    saved = client.post("/api/sales", json={"storage_id": store, "sale_date": _day(0), "lines": [
        {"product_id": line["product_id"], "quantity": line["quantity"], "lot_id": line["lot_id"]} for line in lines]})
    assert saved.status_code == 201, saved.text
    assert sorted((line["quantity"], line["lot_code"] or "") for line in saved.json()["lines"]) == [(2, ""), (7, late["code"])]
    # The named lot gave 7; the unnamed 2 came off the earliest expiry.
    assert [lot["left"] for lot in _lots(client, milk)] == [8, 1]

    too_many = client.post("/api/sales", json={"storage_id": store, "sale_date": _day(0), "lines": [
        {"product_id": milk, "quantity": 2, "lot_id": late["id"]}]})
    assert too_many.status_code == 422 and "only has 1" in too_many.text
    wrong = _upload(client, store, [(milk, None, 1, "L0000-9999")])
    assert wrong.status_code == 422 and "no lot L0000-9999" in wrong.text
