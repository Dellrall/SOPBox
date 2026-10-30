"""Blind stock-count sheets and auditable stocktake adjustments."""

import csv
import io
import zipfile
from datetime import datetime, timezone
from typing import Any

import psycopg
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app import locations
from app.auditlog import audit
from app.errors import ConflictError, InvalidRequestError, NotFoundError
from app.models import StocktakeCreateIn, StocktakeDetail, StocktakeLineOut, StocktakeOut

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_WORKBOOK_BYTES = 25 * 1024 * 1024
MAX_UPLOAD_ROWS = 5000


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def list_products(
    conn: psycopg.Connection, workspace_id: int, storage_id: int | None = None
) -> list[dict[str, Any]]:
    """Active products by category, each with its bundle (the largest pack size, if it has one)
    and its stock in the storage area (the first one when none is given)."""
    storage_id = locations.resolve(conn, workspace_id, storage_id)
    return conn.execute(
        "SELECT p.id, p.name, c.name AS category, p.base_unit, COALESCE(l.quantity, 0) AS stock, "
        "b.unit AS bundle_unit, b.base_qty AS bundle_qty "
        "FROM products p JOIN categories c ON c.id = p.category_id "
        "LEFT JOIN stock_levels l ON l.product_id = p.id AND l.storage_id = %s "
        "LEFT JOIN LATERAL (SELECT unit, base_qty FROM pack_units WHERE product_id = p.id "
        "ORDER BY base_qty DESC, unit LIMIT 1) b ON TRUE "
        "WHERE p.workspace_id = %s AND NOT p.archived ORDER BY lower(c.name), lower(p.name)",
        (storage_id, workspace_id),
    ).fetchall()


HEADERS = ["Product ID", "Product Name", "Bundle Count", "Bundle Unit", "Loose Count", "Base Unit"]
_BAND = PatternFill("solid", fgColor="1F7A35")
_HEAD = PatternFill("solid", fgColor="174C3C")
_NO_BUNDLE = PatternFill("solid", fgColor="E5E7EB")


def build_template(conn: psycopg.Connection, workspace_id: int) -> bytes:
    """A blind-count workbook, one section per category. Products are prefilled, counts are not."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Count sheet"
    instructions = workbook.create_sheet("Instructions")

    for column, width in {"A": 12, "B": 34, "C": 16, "D": 14, "E": 16, "F": 14}.items():
        sheet.column_dimensions[column].width = width
    # Product ID identifies the row on upload; it stays hidden so the sheet reads like the shelf.
    sheet.column_dimensions["A"].hidden = True

    category = None
    for product in list_products(conn, workspace_id):
        if product["category"] != category:
            category = product["category"]
            if sheet.max_row > 1:
                sheet.append([])
            sheet.append([None, category])
            band = sheet.max_row
            sheet.merge_cells(start_row=band, start_column=2, end_row=band, end_column=6)
            title = sheet.cell(band, 2)
            title.data_type = "s"
            title.font = Font(bold=True, color="FFFFFF")
            title.fill = _BAND
            title.alignment = Alignment(horizontal="center")
            sheet.append(HEADERS)
            for cell in sheet[sheet.max_row]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = _HEAD
        sheet.append(
            [product["id"], product["name"], None, product["bundle_unit"], None, product["base_unit"]]
        )
        row = sheet[sheet.max_row]
        # Treat names and units as text even if a value begins with a formula character.
        for cell in (row[1], row[3], row[5]):
            cell.data_type = "s"
        if product["bundle_unit"] is None:
            row[2].fill = _NO_BUNDLE
    if category is None:
        sheet.append(HEADERS)  # no products yet: still a sheet that uploads

    instructions.append(["Stock Count Instructions"])
    instructions.append(["1. Products are grouped by category. Count each one on the shelf."])
    instructions.append(["2. Bundle Count: whole bundles (Bundle Unit). Loose Count: single items (Base Unit)."])
    instructions.append(["3. Extra loose items are fine: 12 loose with 10 per box is saved as 1 box + 2."])
    instructions.append(["4. Leave both counts blank if the product was not counted."])
    instructions.append(["5. Enter 0 only when you counted the product and found none."])
    instructions.append(["6. Do not add rows or move products between sections."])
    instructions.append(["7. Upload this workbook in Stock Count, review the totals, then save the stocktake."])
    instructions.column_dimensions["A"].width = 100
    instructions["A1"].font = Font(bold=True, size=14, color="174C3C")
    for row in range(2, instructions.max_row + 1):
        instructions.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _rows_from_csv(content: bytes) -> list[tuple[int, list[Any]]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise InvalidRequestError("CSV must be UTF-8 encoded") from exc
    reader = csv.reader(io.StringIO(text))
    try:
        rows = list(reader)
    except csv.Error as exc:
        raise InvalidRequestError("the uploaded CSV is not valid") from exc
    if not rows:
        raise InvalidRequestError("the uploaded count sheet is empty")
    return [(index + 1, row) for index, row in enumerate(rows)]


def _rows_from_xlsx(content: bytes) -> list[tuple[int, list[Any]]]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(item.file_size for item in entries) > MAX_WORKBOOK_BYTES:
                raise InvalidRequestError("the workbook is too large to process")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except InvalidRequestError:
        raise
    except (zipfile.BadZipFile, OSError, ValueError, KeyError) as exc:
        raise InvalidRequestError("the uploaded file is not a valid Excel workbook") from exc

    sheet = workbook["Count sheet"] if "Count sheet" in workbook.sheetnames else workbook.active
    rows = [(index, list(row)) for index, row in enumerate(sheet.iter_rows(values_only=True), start=1)]
    workbook.close()
    if not rows:
        raise InvalidRequestError("the uploaded count sheet is empty")
    return rows


def _norm(value: Any) -> str:
    return str(value or "").strip().casefold().replace("_", " ").replace("-", " ")


def _header_indexes(row: list[Any]) -> dict[str, int] | None:
    """Column positions if this row is a header row (it names Product ID), else None.

    A sheet may repeat its header per category section. "Counted Stock" (one total in base
    units) is still read from sheets made before bundle and loose columns existed.
    """
    names = {_norm(value): index for index, value in enumerate(row)}
    found = {
        key: next((names[n] for n in options if n in names), None)
        for key, options in {
            "id": ("product id", "id"),
            "bundles": ("bundle count", "bundles"),
            "loose": ("loose count", "loose"),
            "total": ("counted stock", "counted quantity", "physical count", "quantity"),
        }.items()
    }
    if found["id"] is None:
        return None
    if found["bundles"] is None and found["loose"] is None and found["total"] is None:
        raise InvalidRequestError("the count sheet needs Bundle Count and Loose Count columns")
    return {key: index for key, index in found.items() if index is not None}


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _number(value: Any, *, field: str, row_number: int) -> int:
    try:
        if isinstance(value, bool):
            raise ValueError
        if isinstance(value, float):
            if not value.is_integer():
                raise ValueError
            result = int(value)
        else:
            text = str(value).strip()
            if not text or not text.isdecimal():
                raise ValueError
            result = int(text)
        if result < 0 or result > 9_007_199_254_740_991:
            raise ValueError
        return result
    except (TypeError, ValueError) as exc:
        raise InvalidRequestError(
            f"row {row_number}: {field} must be a whole number of 0 or more"
        ) from exc


def preview_import(
    conn: psycopg.Connection, workspace_id: int, filename: str, content: bytes, storage_id: int | None = None
) -> list[dict[str, Any]]:
    if not content or len(content) > MAX_UPLOAD_BYTES:
        raise InvalidRequestError("upload a non-empty file smaller than 5 MB")
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix == "csv":
        rows = _rows_from_csv(content)
    elif suffix == "xlsx":
        rows = _rows_from_xlsx(content)
    else:
        raise InvalidRequestError("upload an .xlsx or .csv count sheet")
    if len(rows) > MAX_UPLOAD_ROWS:
        raise InvalidRequestError("the count sheet contains too many rows")

    products = {int(p["id"]): p for p in list_products(conn, workspace_id, storage_id)}
    counted: list[dict[str, Any]] = []
    seen: set[int] = set()
    columns: dict[str, int] | None = None
    for row_number, row in rows:
        header = _header_indexes(row)
        if header is not None:
            columns = header
            continue
        if columns is None:
            continue  # a title above the first header row
        cell = lambda key: row[columns[key]] if key in columns and columns[key] < len(row) else None  # noqa: E731
        raw = {key: cell(key) for key in ("bundles", "loose", "total")}
        if all(_blank(value) for value in raw.values()) or _blank(cell("id")):
            continue  # blank is deliberately different from an explicit zero
        product_id = _number(cell("id"), field="Product ID", row_number=row_number)
        if product_id in seen:
            raise InvalidRequestError(f"row {row_number}: product {product_id} appears more than once")
        seen.add(product_id)
        product = products.get(product_id)
        if product is None:
            raise InvalidRequestError(
                f"row {row_number}: product {product_id} is not active in this workspace"
            )
        if not _blank(raw["total"]):
            quantity = _number(raw["total"], field="Counted Stock", row_number=row_number)
        else:
            bundles = 0 if _blank(raw["bundles"]) else _number(raw["bundles"], field="Bundle Count", row_number=row_number)
            loose = 0 if _blank(raw["loose"]) else _number(raw["loose"], field="Loose Count", row_number=row_number)
            if bundles and not product["bundle_qty"]:
                raise InvalidRequestError(
                    f"row {row_number}: {product['name']} has no bundle size; count it in Loose Count"
                )
            quantity = bundles * (product["bundle_qty"] or 0) + loose
        counted.append(
            {
                "product_id": product_id,
                "product_name": product["name"],
                "base_unit": product["base_unit"],
                "expected_stock": int(product["stock"]),
                "counted_stock": quantity,
                "difference": quantity - int(product["stock"]),
            }
        )
    if not counted:
        raise InvalidRequestError("no product counts were entered; blank rows are left unchanged")
    return counted


def _summary(row: dict[str, Any]) -> StocktakeOut:
    return StocktakeOut(
        id=row["id"],
        reference=row["reference"],
        count_date=row["count_date"],
        counted_by=row["counted_by"],
        created_at=row["created_at"],
        notes=row["notes"],
        line_count=row["line_count"],
        changed_count=row["changed_count"],
        storage=row["storage"],
    )


_SUMMARY_SQL = """
SELECT s.id, s.reference, s.count_date, s.counted_by, s.created_at, s.notes,
       COUNT(l.id)::int AS line_count,
       COUNT(l.id) FILTER (WHERE l.difference <> 0)::int AS changed_count,
       (SELECT o.name || ' / ' || st.name FROM storages st JOIN outlets o ON o.id = st.outlet_id
        WHERE st.id = s.storage_id) AS storage
FROM stocktakes s
LEFT JOIN stocktake_lines l ON l.stocktake_id = s.id
"""


def list_stocktakes(conn: psycopg.Connection, workspace_id: int) -> list[StocktakeOut]:
    rows = conn.execute(
        _SUMMARY_SQL
        + " WHERE s.workspace_id = %s GROUP BY s.id "
        "ORDER BY s.count_date DESC, s.id DESC LIMIT 200",
        (workspace_id,),
    ).fetchall()
    return [_summary(row) for row in rows]


def get_stocktake(
    conn: psycopg.Connection, workspace_id: int, stocktake_id: int
) -> StocktakeDetail:
    header = conn.execute(
        _SUMMARY_SQL
        + " WHERE s.workspace_id = %s AND s.id = %s GROUP BY s.id",
        (workspace_id, stocktake_id),
    ).fetchone()
    if header is None:
        raise NotFoundError("stocktake not found")
    rows = conn.execute(
        "SELECT id, product_id, product_name, base_unit, expected_stock, counted_stock, difference "
        "FROM stocktake_lines WHERE stocktake_id = %s ORDER BY lower(product_name)",
        (stocktake_id,),
    ).fetchall()
    return StocktakeDetail(
        **_summary(header).model_dump(),
        lines=[StocktakeLineOut(**row) for row in rows],
    )


def create_stocktake(
    conn: psycopg.Connection,
    workspace_id: int,
    counted_by: str,
    data: StocktakeCreateIn,
) -> StocktakeDetail:
    product_ids = [line.product_id for line in data.lines]
    storage_id = locations.resolve(conn, workspace_id, data.storage_id)
    with conn.transaction():
        # Serialize reference allocation within this workspace and lock product rows before
        # comparing the expected balances submitted by the counter.
        workspace = conn.execute(
            "SELECT id FROM workspaces WHERE id = %s FOR UPDATE", (workspace_id,)
        ).fetchone()
        if workspace is None:
            raise NotFoundError("workspace not found")
        products = conn.execute(
            "SELECT id, name, base_unit FROM products "
            "WHERE workspace_id = %s AND NOT archived AND id = ANY(%s) "
            "ORDER BY id FOR UPDATE",
            (workspace_id, product_ids),
        ).fetchall()
        by_id = {int(product["id"]): product for product in products}
        if set(by_id) != set(product_ids):
            raise InvalidRequestError("one or more products are no longer active")
        have = locations.levels(conn, storage_id, product_ids)
        by_id = {pid: {**product, "stock": have[pid]} for pid, product in by_id.items()}

        for line in data.lines:
            product = by_id[line.product_id]
            if int(product["stock"]) != line.expected_stock:
                raise ConflictError(
                    f"{product['name']} stock changed while the count was open; reload and review it"
                )

        number = conn.execute(
            "SELECT COALESCE(MAX(number), 0) + 1 AS next_number "
            "FROM stocktakes WHERE workspace_id = %s",
            (workspace_id,),
        ).fetchone()["next_number"]
        reference = f"ST-{number:06d}"
        created_at = _now()
        stocktake_id = conn.execute(
            "INSERT INTO stocktakes "
            "(workspace_id, number, reference, count_date, counted_by, created_at, notes, storage_id) "
            "VALUES (%s, %s, %s, %s, %s, %s, %s, %s) RETURNING id",
            (
                workspace_id,
                number,
                reference,
                data.count_date,
                counted_by,
                created_at,
                data.notes.strip(),
                storage_id,
            ),
        ).fetchone()["id"]
        # A count of this storage area does every count planned for it.
        conn.execute(
            "UPDATE count_plans SET stocktake_id = %s WHERE workspace_id = %s AND storage_id = %s AND stocktake_id IS NULL",
            (stocktake_id, workspace_id, storage_id),
        )

        changed_count = 0
        for line in data.lines:
            product = by_id[line.product_id]
            expected = int(product["stock"])
            difference = line.counted_stock - expected
            changed_count += difference != 0
            conn.execute(
                "INSERT INTO stocktake_lines "
                "(stocktake_id, product_id, product_name, base_unit, expected_stock, counted_stock, difference) "
                "VALUES (%s, %s, %s, %s, %s, %s, %s)",
                (
                    stocktake_id,
                    line.product_id,
                    product["name"],
                    product["base_unit"],
                    expected,
                    line.counted_stock,
                    difference,
                ),
            )
            locations.adjust(conn, workspace_id, line.product_id, storage_id, difference)

        audit(
            conn,
            workspace_id,
            counted_by,
            "stocktake.created",
            None,
            f"{reference}: {len(data.lines)} products counted in {locations.storage_label(conn, workspace_id, storage_id)}; "
            f"{changed_count} stock balances changed.",
        )

    return get_stocktake(conn, workspace_id, stocktake_id)
