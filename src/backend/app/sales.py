"""Sales taken off one storage area's stock, from our sales sheet (the count sheet's twin, with the
lot used) or a POS export of codes and quantities.

Taking sales off stock means the next count's difference is only what sales do not explain:
waste, theft or a receiving mistake. More sold than the storage area held takes it to 0, and the
sale keeps both numbers so the gap stays visible.
"""

import csv
import io
from datetime import datetime, timezone
from typing import Any

import psycopg
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

from app import locations, lots, stocktakes
from app.auditlog import audit
from app.errors import InvalidRequestError
from app.models import SaleIn, SaleLineOut, SaleOut, SalePreview, SalePreviewLine

# Header names a sales sheet or POS export may use, compared after stocktakes._norm.
_CODE = ("product code", "code", "sku", "product no", "product number", "item code", "item no", "barcode")
_QUANTITY = ("quantity sold", "qty sold", "sold", "quantity", "qty")


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


# Columns of a sales sheet, compared after stocktakes._norm. Ours names Product ID; a POS export
# names a code and one quantity.
_HEADS = {
    "id": ("product id",),
    "code": _CODE,
    "bundles": ("bundles sold", "bundle sold", "bundle count", "bundles"),
    "loose": ("loose sold", "loose count", "loose"),
    "quantity": _QUANTITY,
    "lot": ("lot used", "lot code", "lot", "label"),
}
SHEET_HEADERS = ["Product ID", "Product Name", "Code", "Bundles Sold", "Bundle Unit", "Loose Sold", "Base Unit", "Lot Used"]


def _columns(row: list[Any]) -> dict[str, int] | None:
    """Column positions if this row is a header row, else None. A sheet may repeat it per category."""
    names = {stocktakes._norm(value): index for index, value in enumerate(row)}
    found = {key: next((names[n] for n in options if n in names), None) for key, options in _HEADS.items()}
    found = {key: index for key, index in found.items() if index is not None}
    if ("id" in found or "code" in found) and found.keys() & {"bundles", "loose", "quantity"}:
        return found
    return None


def _code_text(value: Any) -> str:
    """Excel turns a numeric code into 1001.0; read it back as "1001"."""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def match(
    conn: psycopg.Connection, ws: int, storage_id: int, sold: dict[str, tuple[str, int]]
) -> tuple[list[SalePreviewLine], list[str]]:
    """Each casefolded code's product and quantity, and the codes no active product has."""
    products = {
        r["code"].casefold(): r
        for r in conn.execute(
            "SELECT p.id, p.code, p.name, p.base_unit, COALESCE(l.quantity, 0) AS in_stock FROM products p "
            "LEFT JOIN stock_levels l ON l.product_id = p.id AND l.storage_id = %s "
            "WHERE p.workspace_id = %s AND NOT p.archived AND p.code IS NOT NULL",
            (storage_id, ws),
        ).fetchall()
    }
    lines, unknown = [], []
    for key, (code, quantity) in sold.items():
        product = products.get(key)
        if product is None:
            unknown.append(code)
        elif quantity > 0:
            lines.append(SalePreviewLine(
                product_id=product["id"], code=product["code"], product_name=product["name"],
                base_unit=product["base_unit"], quantity=quantity, in_stock=product["in_stock"],
            ))
    return lines, unknown


def open_lots(conn: psycopg.Connection, ws: int, storage_id: int) -> list[dict]:
    """Lots with something left in this storage area, earliest expiry first, each with its code."""
    rows = [r for r in lots._rows(conn, ws, "AND l.storage_id = %s", (storage_id,)) if r["left_qty"] > 0]
    return [{**r, "code": lots.code(r["number"], r["created_at"])} for r in rows]


def _cell_value(row: list[Any], columns: dict[str, int], key: str) -> Any:
    at = columns.get(key)
    return None if at is None or at >= len(row) or stocktakes._blank(row[at]) else row[at]


def preview(conn: psycopg.Connection, ws: int, storage_id: int, filename: str, content: bytes) -> SalePreview:
    """Find each row's product (by Product ID or code) and lot, add up repeats. Nothing is saved."""
    storage_id = locations.resolve(conn, ws, storage_id)
    if not content or len(content) > stocktakes.MAX_UPLOAD_BYTES:
        raise InvalidRequestError("upload a non-empty file smaller than 5 MB")
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix == "csv":
        rows = stocktakes._rows_from_csv(content)
    elif suffix == "xlsx":
        rows = stocktakes._rows_from_xlsx(content)
    else:
        raise InvalidRequestError("upload an .xlsx or .csv sales sheet")
    if len(rows) > stocktakes.MAX_UPLOAD_ROWS:
        raise InvalidRequestError("the sales sheet contains too many rows")

    products = {int(p["id"]): p for p in stocktakes.list_products(conn, ws, storage_id)}
    code_of = {
        r["id"]: r["code"]
        for r in conn.execute(
            "SELECT id, code FROM products WHERE workspace_id = %s AND NOT archived AND code IS NOT NULL", (ws,)
        ).fetchall()
    }
    codes = {code.casefold(): product_id for product_id, code in code_of.items()}
    lot_codes = {(r["product_id"], r["code"].casefold()): r for r in open_lots(conn, ws, storage_id)}
    columns = None
    sold: dict[tuple[int, int | None], int] = {}  # (product, lot): base units
    unknown: dict[str, str] = {}  # casefolded code: code as first written
    for row_number, row in rows:
        header = _columns(row)
        if header is not None:
            columns = header
            continue
        if columns is None:
            continue  # a title above the header
        raw_id, raw_code = _cell_value(row, columns, "id"), _cell_value(row, columns, "code")
        if raw_id is not None:
            product_id = stocktakes._number(raw_id, field="Product ID", row_number=row_number)
            if product_id not in products:
                raise InvalidRequestError(f"row {row_number}: this product is archived or not in this workspace")
        elif raw_code is not None:
            code = _code_text(raw_code)
            product_id = codes.get(code.casefold())
            if product_id is None:
                unknown.setdefault(code.casefold(), code)
                continue
        else:
            continue
        product = products[product_id]
        bundles, loose, total = (_cell_value(row, columns, key) for key in ("bundles", "loose", "quantity"))
        if bundles is None and loose is None and total is None:
            continue  # not sold
        quantity = 0
        if bundles is not None:
            if not product["bundle_qty"]:
                raise InvalidRequestError(f"row {row_number}: {product['name']} has no bundle size; use Loose Sold")
            quantity += stocktakes._number(bundles, field="Bundles Sold", row_number=row_number) * product["bundle_qty"]
        if loose is not None:
            quantity += stocktakes._number(loose, field="Loose Sold", row_number=row_number)
        if total is not None:
            quantity += stocktakes._number(total, field="Quantity", row_number=row_number)
        if quantity == 0:
            continue
        lot_text = _cell_value(row, columns, "lot")
        lot = None
        if lot_text is not None:
            lot = lot_codes.get((product_id, str(lot_text).strip().casefold()))
            if lot is None:
                raise InvalidRequestError(
                    f"row {row_number}: no lot {str(lot_text).strip()} of {product['name']} is left in this storage area"
                )
        key = (product_id, lot["id"] if lot else None)
        sold[key] = sold.get(key, 0) + quantity

    if columns is None:
        raise InvalidRequestError("the sheet needs a Product code column and a Quantity column")
    by_lot = {r["id"]: r for r in lot_codes.values()}
    lines = [
        SalePreviewLine(
            product_id=product_id, code=code_of.get(product_id, ""), product_name=products[product_id]["name"],
            base_unit=products[product_id]["base_unit"], quantity=quantity, in_stock=products[product_id]["stock"],
            lot_id=lot_id, lot_code=by_lot[lot_id]["code"] if lot_id else None,
        )
        for (product_id, lot_id), quantity in sold.items()
    ]
    if not lines and not unknown:
        raise InvalidRequestError("no sales were found in the sheet")
    lines.sort(key=lambda line: (line.product_name.casefold(), line.lot_code or ""))
    return SalePreview(lines=lines, unknown_codes=sorted(unknown.values()))


def _sales(conn: psycopg.Connection, ws: int, where: str, params: tuple) -> list[SaleOut]:
    rows = conn.execute(
        f"SELECT x.*, {locations.LABEL_SQL} AS label FROM sales x "
        "JOIN storages s ON s.id = x.storage_id JOIN outlets o ON o.id = s.outlet_id "
        f"WHERE x.workspace_id = %s {where} ORDER BY x.sale_date DESC, x.id DESC LIMIT 50",
        (ws, *params),
    ).fetchall()
    lines = conn.execute(
        "SELECT sl.sale_id, sl.product_id, sl.product_code, sl.product_name, sl.base_unit, sl.quantity, sl.deducted, "
        "lot.number AS lot_number, lot.created_at AS lot_created FROM sale_lines sl "
        "LEFT JOIN stock_lots lot ON lot.id = sl.lot_id WHERE sl.sale_id = ANY(%s) ORDER BY lower(sl.product_name), sl.id",
        ([r["id"] for r in rows],),
    ).fetchall()
    return [
        SaleOut(
            id=r["id"], reference=f"SA-{r['number']:06d}", storage=r["label"], sale_date=r["sale_date"],
            source=r["source"], recorded_by=r["recorded_by"], created_at=r["created_at"],
            lines=[
                SaleLineOut(
                    **{k: v for k, v in line.items() if k not in ("sale_id", "lot_number", "lot_created")},
                    lot_code=lots.code(line["lot_number"], line["lot_created"]) if line["lot_number"] else None,
                )
                for line in lines if line["sale_id"] == r["id"]
            ],
        )
        for r in rows
    ]


def list_sales(conn: psycopg.Connection, ws: int) -> list[SaleOut]:
    return _sales(conn, ws, f"AND {locations.mine('x.storage_id')}", ())


def _take_from_lot(conn: psycopg.Connection, ws: int, storage_id: int, product_id: int, lot_id: int, quantity: int) -> None:
    """A sale that names its lot uses that lot instead of the earliest-expiry one, as a disposal does.
    Call before the stock goes down."""
    lot = next((r for r in open_lots(conn, ws, storage_id) if r["id"] == lot_id and r["product_id"] == product_id), None)
    if lot is None:
        raise InvalidRequestError("that lot is not of this product or has nothing left in this storage area")
    if quantity > lot["left_qty"]:
        raise InvalidRequestError(f"lot {lot['code']} only has {lot['left_qty']} {lot['base_unit']} left")
    conn.execute("UPDATE stock_lots SET used = used + %s WHERE id = %s", (quantity, lot_id))


def create_sale(
    conn: psycopg.Connection, ws: int, by: str, data: SaleIn, external_ref: str | None = None
) -> SaleOut:
    """Record a sale. A sale with an external_ref already on record is returned unchanged."""
    storage_id = locations.resolve(conn, ws, data.storage_id)
    product_ids = [line.product_id for line in data.lines]
    with conn.transaction():
        conn.execute("SELECT id FROM workspaces WHERE id = %s FOR UPDATE", (ws,))  # numbers in order
        if external_ref is not None:
            sent = conn.execute(
                "SELECT id FROM sales WHERE workspace_id = %s AND external_ref = %s", (ws, external_ref)
            ).fetchone()
            if sent is not None:
                return _sales(conn, ws, "AND x.id = %s", (sent["id"],))[0]
        products = {
            r["id"]: r
            for r in conn.execute(
                "SELECT id, code, name, base_unit FROM products WHERE workspace_id = %s AND NOT archived "
                "AND id = ANY(%s) ORDER BY id FOR UPDATE",
                (ws, product_ids),
            ).fetchall()
        }
        if set(products) != set(product_ids):
            raise InvalidRequestError("one or more products are unknown or archived")
        have = locations.levels(conn, storage_id, product_ids)
        number = conn.execute(
            "SELECT COALESCE(MAX(number), 0) + 1 AS n FROM sales WHERE workspace_id = %s", (ws,)
        ).fetchone()["n"]
        sale_id = conn.execute(
            "INSERT INTO sales (workspace_id, number, storage_id, sale_date, source, recorded_by, created_at, "
            "external_ref) VALUES (%s, %s, %s, %s, %s, %s, %s, %s) RETURNING id",
            (ws, number, storage_id, data.sale_date, data.source, by, _now(), external_ref),
        ).fetchone()["id"]
        short = 0
        for line in data.lines:
            product = products[line.product_id]
            deducted = min(line.quantity, have[line.product_id])
            have[line.product_id] -= deducted
            short += deducted < line.quantity
            if line.lot_id is not None:
                _take_from_lot(conn, ws, storage_id, line.product_id, line.lot_id, deducted)
            conn.execute(
                "INSERT INTO sale_lines (sale_id, product_id, product_code, product_name, base_unit, quantity, "
                "deducted, lot_id) VALUES (%s, %s, %s, %s, %s, %s, %s, %s)",
                (sale_id, line.product_id, product["code"] or "", product["name"], product["base_unit"],
                 line.quantity, deducted, line.lot_id),
            )
            locations.adjust(conn, ws, line.product_id, storage_id, -deducted)
        audit(
            conn, ws, by, "sales.recorded", None,
            f"SA-{number:06d}: {len(data.lines)} product(s) sold from {locations.storage_label(conn, ws, storage_id)}"
            + (f"; {short} sold more than was in stock" if short else ""),
        )
    return _sales(conn, ws, "AND x.id = %s", (sale_id,))[0]


def _cell(text: str) -> str:
    """Keep a spreadsheet from running a name as a formula."""
    return f"'{text}" if text[:1] in ("=", "+", "-", "@") else text


def build_template(conn: psycopg.Connection, ws: int, storage_id: int | None) -> bytes:
    """The stock count sheet's twin for a day's sales from one storage area. Products that have
    labelled lots there get a Lot Used list, so the sale comes off the box it was taken from."""
    storage_id = locations.resolve(conn, ws, storage_id)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales sheet"
    lot_sheet = workbook.create_sheet("Lots")
    instructions = workbook.create_sheet("Instructions")

    code_of = {
        r["id"]: r["code"]
        for r in conn.execute("SELECT id, code FROM products WHERE workspace_id = %s AND code IS NOT NULL", (ws,)).fetchall()
    }
    lot_sheet.append(["Lot", "Product", "Expiry", "Left", "Base Unit"])
    lot_rows: dict[int, tuple[int, int]] = {}  # product: first and last row on the Lots sheet
    for lot in sorted(open_lots(conn, ws, storage_id), key=lambda r: r["product_id"]):
        lot_sheet.append([lot["code"], lot["product"], lot["expiry_on"], lot["left_qty"], lot["base_unit"]])
        first, _ = lot_rows.get(lot["product_id"], (lot_sheet.max_row, 0))
        lot_rows[lot["product_id"]] = (first, lot_sheet.max_row)
    for cell in lot_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = stocktakes._HEAD
    for column, width in {"A": 14, "B": 34, "C": 12, "D": 10, "E": 12}.items():
        lot_sheet.column_dimensions[column].width = width

    for column, width in {"A": 12, "B": 34, "C": 12, "D": 14, "E": 14, "F": 12, "G": 12, "H": 14}.items():
        sheet.column_dimensions[column].width = width
    sheet.column_dimensions["A"].hidden = True  # Product ID finds the row on upload
    category = None
    for product in stocktakes.list_products(conn, ws, storage_id):
        if product["category"] != category:
            category = product["category"]
            if sheet.max_row > 1:
                sheet.append([])
            sheet.append([None, category])
            band = sheet.max_row
            sheet.merge_cells(start_row=band, start_column=2, end_row=band, end_column=len(SHEET_HEADERS))
            title = sheet.cell(band, 2)
            title.data_type = "s"
            title.font = Font(bold=True, color="FFFFFF")
            title.fill = stocktakes._BAND
            title.alignment = Alignment(horizontal="center")
            sheet.append(SHEET_HEADERS)
            for cell in sheet[sheet.max_row]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = stocktakes._HEAD
        sheet.append([product["id"], _cell(product["name"]), _cell(code_of.get(product["id"], "")), None,
                      product["bundle_unit"], None, product["base_unit"], None])
        row = sheet[sheet.max_row]
        for cell in (row[1], row[2], row[4], row[6]):
            cell.data_type = "s"
        if product["bundle_unit"] is None:
            row[3].fill = stocktakes._NO_BUNDLE
        if product["id"] in lot_rows:
            first, last = lot_rows[product["id"]]
            choice = DataValidation(type="list", formula1=f"Lots!$A${first}:$A${last}", allow_blank=True)
            sheet.add_data_validation(choice)
            choice.add(row[7])
        else:
            row[7].fill = stocktakes._NO_BUNDLE
    if category is None:
        sheet.append(SHEET_HEADERS)

    for line in [
        "Sales Sheet Instructions",
        "1. Products are grouped by category. Fill in only what was sold.",
        "2. Bundles Sold: whole bundles (Bundle Unit). Loose Sold: single items (Base Unit).",
        "3. Lot Used: pick the label on the box the sale came from. Leave it blank and the earliest expiry is used.",
        "4. Sold from two lots? Copy the row below it and put the other lot and amount on the copy.",
        "5. Grey cells do not apply to that product. The Lots tab lists every label with what is left.",
        "6. Upload this workbook in Stock & sales, check the list, then take the sales off stock.",
    ]:
        instructions.append([line])
    instructions.column_dimensions["A"].width = 100
    instructions["A1"].font = Font(bold=True, size=14, color="174C3C")
    for row in range(2, instructions.max_row + 1):
        instructions.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def template_csv(conn: psycopg.Connection, ws: int) -> bytes:
    """Every active product that has a code, with Quantity left for the day's sales."""
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(["Product code", "Product name", "Quantity"])
    for r in conn.execute(
        "SELECT code, name FROM products WHERE workspace_id = %s AND NOT archived AND code IS NOT NULL "
        "ORDER BY lower(name)",
        (ws,),
    ).fetchall():
        writer.writerow([_cell(r["code"]), _cell(r["name"]), ""])
    return out.getvalue().encode("utf-8-sig")
